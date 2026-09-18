"""Build clean pre-submission DOCX files from the reviewed Markdown sources."""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"


def set_font(run, name="Arial", size=10.5, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, bold=True)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, italic=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, name="Consolas", size=9)
        else:
            set_font(paragraph.add_run(part))


def configure(doc, line_numbering=False):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    if line_numbering:
        line_numbers = OxmlElement("w:lnNumType")
        line_numbers.set(qn("w:countBy"), "5")
        line_numbers.set(qn("w:restart"), "continuous")
        sec._sectPr.append(line_numbers)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = footer.add_run()
    run._r.extend([begin, instruction, end])
    set_font(run, size=9)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = None


def build(md_path, out_path, include_figures=False):
    doc = Document()
    configure(doc, line_numbering=include_figures)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    in_code = False
    figure_inserted = set()
    for line in lines:
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph()
            set_font(p.add_run(line), name="Consolas", size=8.5)
            continue
        if not line.strip() or line.strip() == "---":
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(line[2:].strip()), size=18, bold=True)
            continue
        heading = re.match(r"^(#{2,4})\s+(.*)$", line)
        if heading:
            level = min(len(heading.group(1)) - 1, 3)
            doc.add_heading(heading.group(2), level=level)
            continue
        fig = re.match(r"^\*\*Figure ([1-6])\*\*", line)
        if include_figures and fig and fig.group(1) not in figure_inserted:
            names = {
                "1": "fig1_workflow.png", "2": "fig2_target_evidence.png",
                "3": "fig3_screening.png", "4": "fig4_bound_md_v2.png",
                "5": "fig5_models_synthesis.png", "6": "fig6_audit_trajectory.png",
            }
            image = PAPER / "figures" / names[fig.group(1)]
            if image.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                shape = p.add_run().add_picture(str(image), width=Inches(6.4))
                doc_pr = shape._inline.docPr
                doc_pr.set("title", f"Figure {fig.group(1)}")
                doc_pr.set("descr", f"Figure {fig.group(1)}; full description follows in the figure legend.")
                figure_inserted.add(fig.group(1))
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", line))
        elif line.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote")
            add_inline(p, line[2:])
        elif line.startswith("|"):
            # Markdown tables remain readable as compact monospaced rows.
            if set(line.replace("|", "").replace("-", "").replace(":", "").strip()):
                p = doc.add_paragraph()
                set_font(p.add_run(line.strip("|").replace("|", "  |  ")), name="Consolas", size=8)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.08
            add_inline(p, line)
    doc.save(out_path)


def main():
    outputs = [
        (PAPER / "SR_MANUSCRIPT_PRESUBMISSION_v2.md", PAPER / "SR_MANUSCRIPT_PRESUBMISSION_v2.docx", True),
        (PAPER / "cover_letter_presubmission_v2.md", PAPER / "cover_letter_presubmission_v2.docx", False),
        (PAPER / "supporting_information_presubmission_v2.md", PAPER / "Supporting_Information_presubmission_v2.docx", False),
    ]
    for source, target, figures in outputs:
        build(source, target, figures)
        print(target)


if __name__ == "__main__":
    sys.exit(main())
